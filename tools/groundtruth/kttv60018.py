from __future__ import annotations

import argparse
import json
import math
import re
import unicodedata
import urllib.parse
import urllib.request
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from typing import Any

from common import atomic_write, build_observation, sha256_bytes, sha256_file, utc_now_iso, write_json, write_jsonl

try:
    from openpyxl import load_workbook
except ImportError as exc:
    raise SystemExit("openpyxl is required: python -m pip install openpyxl") from exc

SOURCE_NAME = "KTTV_AUTO_60018_XLSX"
SOURCE_URL = "https://kttvtudong.net/kttv/export/excelexport?sid=33"
NAMESPACE = "KTTV_AUTO"
IDENTIFIER = "60018"
STATION_EPOCH = "SID33_PUBLIC_AUTOMATED_STATION"

FIELD_MAP = {
    "ff": ("wind_speed", "m/s", "numeric"),
    "dd": ("wind_direction", "deg", "numeric"),
    "fxfx2m": ("wind_gust_2m", "m/s", "numeric"),
    "dxdx2m": ("wind_gust_direction_2m", "deg", "numeric"),
    "tgxh2m": ("wind_gust_time_2m", None, "text"),
    "fxfx2s": ("wind_gust_2s", "m/s", "numeric"),
    "dxdx2s": ("wind_gust_direction_2s", "deg", "numeric"),
    "tgxh2s": ("wind_gust_time_2s", None, "text"),
    "h": ("wave_height_h", "m", "numeric"),
    "tm02": ("wave_period_tm02", "s", "numeric"),
    "hmax": ("wave_height_max", "m", "numeric"),
    "hm0": ("significant_wave_height_hm0", "m", "numeric"),
}
TIME_ALIASES = {"thoigian", "ngaygio", "datetime", "time", "timestamp", "ngay"}


def norm_text(value: Any) -> str:
    if value is None:
        return ""
    s = unicodedata.normalize("NFKD", str(value))
    s = "".join(ch for ch in s if not unicodedata.combining(ch)).lower().strip()
    return re.sub(r"[^a-z0-9]+", "", s)


def _to_iso(value: Any) -> str | None:
    tz = timezone(timedelta(hours=7))
    if isinstance(value, datetime):
        return (value.replace(tzinfo=tz) if value.tzinfo is None else value).isoformat()
    if isinstance(value, date):
        return datetime.combine(value, time.min).replace(tzinfo=tz).isoformat()
    if isinstance(value, str):
        s = value.strip()
        if not s:
            return None
        fmts = [
            "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M",
            "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M",
        ]
        for fmt in fmts:
            try:
                return datetime.strptime(s, fmt).replace(tzinfo=tz).isoformat()
            except ValueError:
                pass
    return None


def _numeric(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and math.isnan(value):
            return None
        return float(value)
    if isinstance(value, str):
        s = value.strip().replace(",", ".")
        if not s or s in {"-", "--", "/", "//", "NA", "N/A"}:
            return None
        try:
            return float(s)
        except ValueError:
            return None
    return None


def discover_layout(ws) -> dict[str, Any]:
    candidates: list[dict[str, Any]] = []
    scan_rows = min(ws.max_row, 25)
    for ridx in range(1, scan_rows + 1):
        row = [ws.cell(ridx, c).value for c in range(1, ws.max_column + 1)]
        normalized = [norm_text(v) for v in row]
        fields: dict[str, int] = {}
        time_col = None
        for cidx, key in enumerate(normalized, start=1):
            if key in FIELD_MAP:
                fields[key] = cidx
            if key in TIME_ALIASES:
                time_col = cidx
        score = len(fields)
        if score:
            candidates.append({"row": ridx, "score": score, "fields": fields, "time_col": time_col, "raw": row})
    if not candidates:
        raise ValueError(f"Could not find 60018 header row in sheet {ws.title!r}")
    best = max(candidates, key=lambda x: (x["score"], -x["row"]))
    if best["score"] < 3:
        raise ValueError(f"Header confidence too low: only {best['score']} known fields")

    if best["time_col"] is None:
        best_count = 0
        best_col = None
        for cidx in range(1, ws.max_column + 1):
            count = 0
            for ridx in range(best["row"] + 1, min(ws.max_row, best["row"] + 30) + 1):
                if _to_iso(ws.cell(ridx, cidx).value):
                    count += 1
            if count > best_count:
                best_count, best_col = count, cidx
        if best_count >= 2:
            best["time_col"] = best_col
    if best["time_col"] is None:
        raise ValueError("Could not resolve observation timestamp column")
    return best


def pick_sheet(wb):
    scored = []
    for ws in wb.worksheets:
        try:
            layout = discover_layout(ws)
            scored.append((layout["score"], ws, layout))
        except ValueError:
            continue
    if not scored:
        raise ValueError("No worksheet contains the expected 60018 field schema")
    scored.sort(key=lambda x: x[0], reverse=True)
    return scored[0][1], scored[0][2]


def parse_workbook(
    src: str | Path,
    dst: str | Path,
    *,
    source_url: str = SOURCE_URL,
    fetched_at: str | None = None,
    raw_path: str | None = None,
    raw_sha256: str | None = None,
) -> dict[str, Any]:
    src = Path(src)
    manifest_path = Path(str(src) + ".manifest.json")
    if manifest_path.exists():
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        source_url = source_url or manifest.get("source_url")
        fetched_at = fetched_at or manifest.get("fetched_at")
        raw_sha256 = raw_sha256 or manifest.get("sha256")
        raw_path = raw_path or manifest.get("raw_path")
    raw_sha256 = raw_sha256 or sha256_file(src)
    raw_path = raw_path or str(src)
    wb = load_workbook(src, read_only=True, data_only=True)
    ws, layout = pick_sheet(wb)

    observations = []
    rows_with_time = 0
    rows_with_values = 0
    missing_time_rows = 0
    first_time = None
    last_time = None
    for ridx in range(layout["row"] + 1, ws.max_row + 1):
        source_time = _to_iso(ws.cell(ridx, layout["time_col"]).value)
        present = []
        for field, cidx in layout["fields"].items():
            val = ws.cell(ridx, cidx).value
            if val is not None and str(val).strip() != "":
                present.append((field, val))
        if not present:
            continue
        if not source_time:
            missing_time_rows += 1
            continue
        rows_with_time += 1
        rows_with_values += 1
        first_time = first_time or source_time
        last_time = source_time
        for field, raw_value in present:
            metric, unit, kind = FIELD_MAP[field]
            if kind == "numeric":
                value = _numeric(raw_value)
                if value is None:
                    continue
            else:
                value = str(raw_value).strip()
            flags = ["UNIT_FROM_60018_FIELD_SCHEMA"] if unit else []
            observations.append(build_observation(
                namespace=NAMESPACE,
                identifier=IDENTIFIER,
                station_epoch=STATION_EPOCH,
                source_time=source_time,
                metric=metric,
                value=value,
                unit=unit,
                source=SOURCE_NAME,
                source_url=source_url,
                fetched_at=fetched_at,
                raw_sha256=raw_sha256,
                raw_path=raw_path,
                raw_report=None,
                raw_group=f"{field}={raw_value}",
                observation_type="INSTRUMENT_AUTOMATED",
                qc_status="PASS",
                qc_flags=flags,
                extra={
                    "kttv60018": {
                        "sheet": ws.title,
                        "row": ridx,
                        "field": field,
                        "header_row": layout["row"],
                    }
                },
            ))
    count = write_jsonl(dst, observations)
    return {
        "sheet": ws.title,
        "header_row": layout["row"],
        "time_col": layout["time_col"],
        "fields": layout["fields"],
        "rows_with_time": rows_with_time,
        "rows_with_values": rows_with_values,
        "rows_missing_time": missing_time_rows,
        "first_time": first_time,
        "last_time": last_time,
        "observations": count,
        "metrics": dict(sorted(__import__("collections").Counter(o["metric"] for o in observations).items())),
        "raw_sha256": raw_sha256,
    }


def fetch_raw(fd: str, td: str, archive_root: str | Path) -> dict[str, Any]:
    for label, value in (("fd", fd), ("td", td)):
        try:
            datetime.strptime(value, "%d/%m/%Y")
        except ValueError as exc:
            raise ValueError(f"{label} must be DD/MM/YYYY") from exc
    data = urllib.parse.urlencode({"fd": fd, "td": td}).encode()
    req = urllib.request.Request(SOURCE_URL, data=data, method="POST", headers={
        "User-Agent": "Mozilla/5.0 (compatible; JoTrip-groundtruth-ingest/1.0)",
        "Referer": "https://kttvtudong.net/kttv/detail/view?sid=33",
        "Content-Type": "application/x-www-form-urlencoded; charset=UTF-8",
        "Accept": "*/*",
    })
    fetched_at = utc_now_iso()
    with urllib.request.urlopen(req, timeout=45) as resp:
        blob = resp.read()
        status = resp.status
        ctype = resp.headers.get("Content-Type")
        disposition = resp.headers.get("Content-Disposition")
    if status != 200:
        raise RuntimeError(f"KTTV returned HTTP {status}")
    if not blob.startswith(b"PK"):
        raise RuntimeError("KTTV response is not an XLSX/ZIP payload")
    digest = sha256_bytes(blob)
    start = datetime.strptime(fd, "%d/%m/%Y")
    end = datetime.strptime(td, "%d/%m/%Y")
    root = Path(archive_root) / "kttv60018" / start.strftime("%Y") / start.strftime("%m")
    name = f"kttv60018_{start:%Y%m%d}_{end:%Y%m%d}_{digest[:12]}.xlsx"
    path = root / name
    if path.exists():
        existing = sha256_file(path)
        if existing != digest:
            raise RuntimeError(f"Immutable archive collision at {path}")
    else:
        atomic_write(path, blob)
    manifest = {
        "schema_version": "jotrip-groundtruth-raw-manifest-v1",
        "source": SOURCE_NAME,
        "source_url": SOURCE_URL,
        "namespace": NAMESPACE,
        "identifier": IDENTIFIER,
        "station_epoch": STATION_EPOCH,
        "fd": fd,
        "td": td,
        "fetched_at": fetched_at,
        "http_status": status,
        "content_type": ctype,
        "content_disposition": disposition,
        "bytes": len(blob),
        "sha256": digest,
        "raw_path": str(path),
    }
    write_json(str(path) + ".manifest.json", manifest)
    return manifest


def main() -> None:
    ap = argparse.ArgumentParser(description="JoTrip KTTV 60018 groundtruth ingestion")
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("normalize")
    p.add_argument("--input", required=True)
    p.add_argument("--output", required=True)
    p.add_argument("--source-url", default=SOURCE_URL)
    p.add_argument("--fetched-at")
    p.add_argument("--raw-path")
    p.add_argument("--raw-sha256")
    f = sub.add_parser("fetch")
    f.add_argument("--from", dest="fd", required=True)
    f.add_argument("--to", dest="td", required=True)
    f.add_argument("--archive-root", default="data/groundtruth/raw")
    args = ap.parse_args()
    if args.cmd == "fetch":
        print(json.dumps(fetch_raw(args.fd, args.td, args.archive_root), ensure_ascii=False, indent=2))
    else:
        print(json.dumps(parse_workbook(
            args.input, args.output, source_url=args.source_url, fetched_at=args.fetched_at,
            raw_path=args.raw_path, raw_sha256=args.raw_sha256,
        ), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
